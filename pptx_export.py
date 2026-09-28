"""Exporta la presentación a PowerPoint (.pptx) sobre una plantilla, con fórmulas editables.

La fuente de verdad sigue siendo Beamer: se exportan los frames que el pipeline ya validó y
compiló. Cada frame se descompone en piezas (texto y listas, bloques, columnas, tablas,
figuras); python-pptx las ubica sobre la plantilla y pandoc traduce el texto LaTeX (negritas,
código y sobre todo las fórmulas, que quedan como ecuaciones nativas OMML de Office) a
párrafos de PowerPoint. Lo que PowerPoint no puede representar (pseudocódigo algorithm2e,
tikz) se inserta como imagen recortada de la diapositiva Beamer compilada.

La plantilla (por defecto plantilla.pptx) aporta el maestro (logos, franjas, pie, número de
diapositiva) y su primera diapositiva es la portada: se reemplazan sus textos.
"""
from __future__ import annotations

import copy
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

DEFAULT_PLANTILLA = Path(__file__).with_name("plantilla.pptx")

# Estilo de la plantilla por defecto (Escuela de Ingeniería Informática PUCV)
COLOR_TITULO = RGBColor(0x95, 0x37, 0x35)      # accent2 al 75 %
COLOR_TEXTO = RGBColor(0x00, 0x20, 0x60)
COLOR_ALERT = COLOR_TITULO
COLOR_CAJA = {"block": RGBColor(0x4F, 0x81, 0xBD), "alertblock": COLOR_TITULO,
              "exampleblock": RGBColor(0x1F, 0x49, 0x7D)}
FONDO_CAJA = RGBColor(0xDC, 0xE6, 0xF2)
FUENTE_TITULO = "Roboto Medium"

# Zonas de la diapositiva de contenido (EMU, plantilla 4:3)
TITULO_XY = (Emu(770400), Emu(1000000), Emu(7800000), Emu(560000))
CUERPO = (Emu(573887), Emu(1640000), Emu(8000000), Emu(4780000))   # x, y, ancho, alto
TAMANOS = (20, 18, 16, 14, 13, 12, 11)
EMU_PT = 12700

ENTORNOS = ("columns", "block", "alertblock", "exampleblock", "tabular", "tabular*", "center",
            "algorithm", "tikzpicture", "figure", "table")
RASTER = ("algorithm", "tikzpicture")


def _pandoc() -> str:
    p = shutil.which("pandoc")
    if p:
        return p
    import pypandoc                                  # pypandoc_binary trae pandoc
    return pypandoc.get_pandoc_path()


# --------------------------------------------------------------------------
# LaTeX del frame → piezas
# --------------------------------------------------------------------------
def _cierre(s: str, i: int, env: str) -> int:
    """Índice del \\end{env} que cierra el \\begin{env} que empieza antes de i."""
    abre, cierra = f"\\begin{{{env}}}", f"\\end{{{env}}}"
    depth, j = 1, i
    while depth:
        a, c = s.find(abre, j), s.find(cierra, j)
        if c < 0:
            return len(s)
        if 0 <= a < c:
            depth, j = depth + 1, a + len(abre)
        else:
            depth, j = depth - 1, c + len(cierra)
    return j - len(cierra)


def _arg(s: str, i: int) -> tuple[str, int]:
    """Argumento {…} (con llaves anidadas) que empieza en s[i]; devuelve (texto, fin)."""
    while i < len(s) and s[i] in " \n\t":
        i += 1
    if i >= len(s) or s[i] != "{":
        return "", i
    depth = 0
    for j in range(i, len(s)):
        depth += {"{": 1, "}": -1}.get(s[j], 0)
        if depth == 0:
            return s[i + 1:j], j + 1
    return s[i + 1:], len(s)


def _opt(s: str, i: int) -> tuple[str, int]:
    while i < len(s) and s[i] in " \n\t":
        i += 1
    if i < len(s) and s[i] == "[":
        j = s.index("]", i)
        return s[i + 1:j], j + 1
    return "", i


def partes_frame(frame: str) -> tuple[str, str]:
    """(título, cuerpo) de \\begin{frame}[…]{título} … \\end{frame}."""
    m = re.search(r"\\begin\{frame\}", frame)
    i = m.end() if m else 0
    _, i = _opt(frame, i)
    titulo, i = _arg(frame, i)
    fin = frame.rfind("\\end{frame}")
    return titulo.strip(), frame[i:fin if fin >= 0 else len(frame)]


def piezas(cuerpo: str) -> list[dict]:
    """Descompone el cuerpo en piezas ordenadas: texto, bloque, columnas, tabla, imagen, raster."""
    cuerpo = re.sub(r"(?<!\\)%.*", "", cuerpo)
    out, i, texto = [], 0, []
    patron = re.compile(r"\\begin\{(" + "|".join(re.escape(e) for e in ENTORNOS) + r")\}|\\includegraphics")

    def vaciar():
        t = "".join(texto).strip()
        if re.sub(r"\\(small|footnotesize|scriptsize|tiny|normalsize|centering|medskip|smallskip|"
                  r"bigskip|vfill|hfill|vspace\*?\{[^}]*\}|pause)", "", t).strip():
            out.append({"t": "texto", "latex": t})
        texto.clear()

    while (m := patron.search(cuerpo, i)):
        texto.append(cuerpo[i:m.start()])
        if m.group(0) == "\\includegraphics":
            _, j = _opt(cuerpo, m.end())
            ruta, j = _arg(cuerpo, j)
            vaciar()
            out.append({"t": "imagen", "ruta": ruta.strip()})
            i = j
            continue
        env = m.group(1)
        j = m.end()
        fin = _cierre(cuerpo, j, env)
        dentro = cuerpo[j:fin]
        i = fin + len(f"\\end{{{env}}}")
        vaciar()
        if env in ("block", "alertblock", "exampleblock"):
            titulo, k = _arg(dentro, 0)
            out.append({"t": "bloque", "clase": env, "titulo": titulo, "hijos": piezas(dentro[k:])})
        elif env == "columns":
            _, k = _opt(dentro, 0)
            cols = re.split(r"\\column\s*(?:\[[^\]]*\])?\s*\{([^}]*)\}|\\begin\{column\}\s*(?:\[[^\]]*\])?"
                            r"\s*\{([^}]*)\}", dentro[k:])
            lista, anchos = [], []
            for n in range(1, len(cols), 3):
                ancho = cols[n] or cols[n + 1] or "0.48\\textwidth"
                f = re.match(r"\s*([\d.]+)", ancho)
                anchos.append(float(f.group(1)) if f else 0.48)
                lista.append(piezas(re.sub(r"\\end\{column\}", "", cols[n + 2])))
            out.append({"t": "columnas", "anchos": anchos, "cols": lista})
        elif env.startswith("tabular"):
            out.append({"t": "tabla", "filas": _filas(dentro, env)})
        elif env in RASTER:
            out.append({"t": "raster", "entorno": env})
        else:                                             # center, figure, table: su contenido
            out += piezas(re.sub(r"\\caption\{[^}]*\}", "", dentro))
    texto.append(cuerpo[i:])
    vaciar()
    return out


def _filas(dentro: str, env: str) -> list[list[str]]:
    k = 0
    if env == "tabular*":
        _, k = _arg(dentro, k)
    _, k = _arg(dentro, k)                                # especificación de columnas
    cuerpo = re.sub(r"\\(toprule|midrule|bottomrule|hline|cline\{[^}]*\}|cmidrule(\([^)]*\))?\{[^}]*\})",
                    "", dentro[k:])
    filas = []
    for fila in re.split(r"\\\\(?:\[[^\]]*\])?", cuerpo):
        if not fila.strip():
            continue
        filas.append([c.strip() for c in re.split(r"(?<!\\)&", fila)])
    return filas


# --------------------------------------------------------------------------
# Texto LaTeX → párrafos de PowerPoint (pandoc, en una sola llamada por presentación)
# --------------------------------------------------------------------------
LIMPIAR = [(r"\\(small|footnotesize|scriptsize|tiny|normalsize|large|Large|centering|medskip|"
            r"smallskip|bigskip|vfill|hfill|pause|noindent)\b", ""),
           (r"\\vspace\*?\{[^}]*\}|\\hspace\*?\{[^}]*\}", " "),
           (r"\\alert<[^>]*>", r"\\alert"), (r"\\and\b", ", "), (r"\\titlepage|\\tableofcontents", "")]


def _preparar(t: str) -> str:
    for a, b in LIMPIAR:
        t = re.sub(a, b, t)
    return t


def a_parrafos(textos: list[str], macros: str = "") -> list[list]:
    """Cada texto LaTeX → lista de <a:p> (lxml). \\alert se marca subrayado y luego se pinta."""
    if not textos:
        return []
    doc = [macros, r"\newcommand{\alert}[1]{\underline{#1}}"]
    for n, t in enumerate(textos):
        doc.append(f"\\section{{PIEZA{n}}}\n{_preparar(t)}\n")
    with tempfile.TemporaryDirectory() as d:
        src, dst = Path(d) / "p.tex", Path(d) / "p.pptx"
        src.write_text("\n".join(doc))
        r = subprocess.run([_pandoc(), "-f", "latex", "-t", "pptx", str(src), "-o", str(dst)],
                           capture_output=True, text=True, timeout=120)
        if r.returncode:
            raise RuntimeError(f"pandoc no pudo leer el texto de las diapositivas: {r.stderr[-600:]}")
        prs = Presentation(str(dst))
    res: dict[int, list] = {}
    actual = None
    for s in prs.slides:
        # se recorre el XML: con fórmulas, pandoc envuelve la caja en mc:AlternateContent
        for sp in s.shapes._spTree.iter(qn("p:sp")):
            ph = sp.find(".//" + qn("p:ph"))
            ps = sp.findall(".//" + qn("a:p"))
            texto = "".join(t.text or "" for t in sp.iter(qn("a:t"))).strip()
            if ph is not None and ph.get("type") in ("title", "ctrTitle"):
                m = re.fullmatch(r"PIEZA(\d+)", texto)
                if m:
                    actual = int(m.group(1))
                    res.setdefault(actual, [])
                continue
            if actual is not None:
                res[actual] += [copy.deepcopy(q) for q in ps]
    return [res.get(n, []) for n in range(len(textos))]


def _texto_plano(t: str) -> str:
    t = re.sub(r"\\(texttt|textbf|emph|alert|textit|mathrm|operatorname)\{([^{}]*)\}", r"\2", t)
    t = re.sub(r"\\[a-zA-Z]+\*?(\[[^\]]*\])?", " ", t)
    return " ".join(re.sub(r"[{}$\\]", "", t).split())


# --------------------------------------------------------------------------
# Medidas (estimación: sin motor de texto, por caracteres y líneas)
# --------------------------------------------------------------------------
def _alto_texto(latex: str, ancho: int, pt: int) -> int:
    lineas = 0.0
    por_linea = max(8, ancho / (pt * EMU_PT * 0.58))
    partes = re.split(r"\\item|\\\\|\n\s*\n", latex)
    for p in partes:
        disp = re.findall(r"\\\[(.*?)\\\]|\\begin\{(?:equation|align)\*?\}(.*?)\\end", p, re.S)
        for d in disp:
            lineas += 2.4 if "\\frac" in "".join(d) or "\\sum" in "".join(d) else 1.6
        p = re.sub(r"\\\[.*?\\\]|\\begin\{(equation|align)\*?\}.*?\\end\{\1\*?\}", "", p, flags=re.S)
        txt = _texto_plano(p)
        if txt:
            lineas += max(1, -(-len(txt) // int(por_linea)))
    return int(lineas * pt * EMU_PT * 1.3) + int(pt * EMU_PT * 0.4)


def _alto(pz: dict, ancho: int, pt: int) -> int:
    t = pz["t"]
    if t == "texto":
        return _alto_texto(pz["latex"], ancho, pt)
    if t == "bloque":
        return int(pt * EMU_PT * 1.9) + sum(_alto(h, ancho - 2 * 91440, pt) for h in pz["hijos"]) + 91440
    if t == "columnas":
        tot = sum(pz["anchos"]) or 1
        return max((sum(_alto(h, int(ancho * a / tot) - 91440, pt) for h in col)
                    for a, col in zip(pz["anchos"], pz["cols"])), default=0)
    if t == "tabla":
        return len(pz["filas"]) * int(min(pt, 14) * EMU_PT * 1.9) + 91440
    return int(CUERPO[3] * 0.62)                       # imagen o raster: se ajusta después


def _hay_imagen(pzs: list[dict]) -> bool:
    return any(p["t"] in ("imagen", "raster") or (p["t"] == "bloque" and _hay_imagen(p["hijos"]))
               or (p["t"] == "columnas" and any(_hay_imagen(c) for c in p["cols"])) for p in pzs)


# --------------------------------------------------------------------------
# Dibujo
# --------------------------------------------------------------------------
def _caja(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for lado in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, lado, Emu(45720))
    return tb, tf


def _pintar(p, pt: int, color: RGBColor, centrar: bool = False):
    """Tamaño y color en cada run; \\alert (subrayado) → color de alerta en negrita."""
    ppr = p.find(qn("a:pPr"))
    if ppr is None:
        ppr = etree.SubElement(p, qn("a:pPr"))
        p.remove(ppr)
        p.insert(0, ppr)
    if ppr.get("lvl") is not None and ppr.find(qn("a:buNone")) is None:   # ítem de lista
        lvl = int(ppr.get("lvl"))
        ppr.set("marL", str(int((lvl + 1) * 0.28 * 914400)))
        ppr.set("indent", str(-int(0.22 * 914400)))
        for b in list(ppr):
            if b.tag.startswith("{%s}bu" % ppr.nsmap.get("a", "")):
                ppr.remove(b)
        bu = etree.SubElement(ppr, qn("a:buChar"))
        bu.set("char", "•" if lvl == 0 else "–")
    if centrar:
        ppr.set("algn", "ctr")
    for r in p.iter(qn("a:rPr"), qn("a:endParaRPr")):
        r.set("sz", str(pt * 100))
        alert = r.get("u") == "sng"
        if alert:
            del r.attrib["u"]
            r.set("b", "1")
        for f in r.findall(qn("a:solidFill")):
            r.remove(f)
        sf = etree.Element(qn("a:solidFill"))
        c = etree.SubElement(sf, qn("a:srgbClr"))
        c.set("val", str(COLOR_ALERT if alert else color))
        r.insert(0, sf)
    return p


def _parrafos_en(tf, parrafos: list, pt: int, color: RGBColor, primero: bool = True):
    txbody = tf._txBody
    if primero:
        for p in txbody.findall(qn("a:p")):
            txbody.remove(p)
    for p in parrafos:
        es_disp = p.find(".//{http://schemas.openxmlformats.org/drawingml/2010/main}m") is not None \
            and not [t for t in p.iter(qn("a:t")) if t.text and t.text.strip()]
        txbody.append(_pintar(copy.deepcopy(p), pt, color, centrar=es_disp))
    if not txbody.findall(qn("a:p")):
        etree.SubElement(txbody, qn("a:p"))


class Dibujante:
    def __init__(self, prs, parrafos: dict[int, list], figuras: Path | None, raster, fuentes: dict):
        self.prs, self.par, self.figuras, self.raster = prs, parrafos, figuras, raster
        self.fuentes = fuentes         # id del elemento de la caja → (latex, pt, color) para el fallback

    def pieza(self, slide, pz: dict, x, y, w, pt, alto_libre, reservado: int = 0) -> int:
        t = pz["t"]
        if t == "texto":
            h = _alto_texto(pz["latex"], w, pt)
            caja, tf = _caja(slide, x, y, w, h)
            _parrafos_en(tf, self.par[pz["id"]], pt, COLOR_TEXTO)
            self.fuentes[caja.shape_id] = (pz["latex"], pt, COLOR_TEXTO, False)
            return h
        if t == "bloque":
            hh = int(pt * EMU_PT * 1.8)
            cab, tf = _caja(slide, x, y, w, hh)
            cab.fill.solid()
            cab.fill.fore_color.rgb = COLOR_CAJA[pz["clase"]]
            _parrafos_en(tf, self.par[pz["id_titulo"]], pt, RGBColor(0xFF, 0xFF, 0xFF))
            self.fuentes[cab.shape_id] = (r"\textbf{%s}" % pz["titulo"], pt, RGBColor(0xFF, 0xFF, 0xFF), True)
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.bold = True
            cuerpo_h = sum(_alto(h, w - 182880, pt) for h in pz["hijos"]) + 91440
            fondo = slide.shapes.add_shape(1, x, y + hh, w, cuerpo_h)
            fondo.fill.solid()
            fondo.fill.fore_color.rgb = FONDO_CAJA
            fondo.line.fill.background()
            yy = y + hh + 45720
            for hijo in pz["hijos"]:
                yy += self.pieza(slide, hijo, x + 91440, yy, w - 182880, pt, alto_libre - (yy - y))
            return hh + cuerpo_h + 91440
        if t == "columnas":
            tot, xx, alto = sum(pz["anchos"]) or 1, x, 0
            sep = 137160
            for a, col in zip(pz["anchos"], pz["cols"]):
                cw = int((w - sep * (len(pz["cols"]) - 1)) * a / tot)
                yy = y
                for k, hijo in enumerate(col):
                    resto = sum(_alto(h, cw, pt) for h in col[k + 1:] if h["t"] not in ("imagen", "raster"))
                    yy += self.pieza(slide, hijo, xx, yy, cw, pt, alto_libre - (yy - y), resto)
                alto, xx = max(alto, yy - y), xx + cw + sep
            return alto
        if t == "tabla":
            return self.tabla(slide, pz, x, y, w, min(pt, 14))
        if t in ("imagen", "raster"):
            png = self.imagen(pz)
            if not png:
                return 0
            import pymupdf
            pm = pymupdf.Pixmap(str(png))
            iw, ih = pm.width, pm.height
            h_max = max(int(alto_libre - reservado), 914400)
            esc = min(w / iw, h_max / ih)
            ancho, alto = int(iw * esc), int(ih * esc)
            slide.shapes.add_picture(str(png), x + (w - ancho) // 2, y, ancho, alto)
            return alto + 45720
        return 0

    def imagen(self, pz):
        if pz["t"] == "imagen":
            if self.figuras is None:
                return None
            p = self.figuras.parent / pz["ruta"] if pz["ruta"].startswith("figuras/") else self.figuras / pz["ruta"]
            if p.suffix == ".pdf":                    # diagramas: PDF en Beamer, PNG en PowerPoint
                p = p.with_suffix(".png")
            return p if p.exists() else None
        return self.raster(pz)

    def tabla(self, slide, pz, x, y, w, pt) -> int:
        filas = pz["filas"]
        ncol = max(len(f) for f in filas)
        fh = int(pt * EMU_PT * 1.9)
        largos = [max((len(_texto_plano(f[c])) if c < len(f) else 0) for f in filas) + 3 for c in range(ncol)]
        forma = slide.shapes.add_table(len(filas), ncol, x, y, w, fh * len(filas))
        tabla = forma.table
        for c in range(ncol):
            tabla.columns[c].width = int(w * largos[c] / sum(largos))
        for r, fila in enumerate(filas):
            tabla.rows[r].height = fh
            for c in range(ncol):
                celda = tabla.cell(r, c)
                celda.margin_top = celda.margin_bottom = Emu(18288)
                ids = pz["ids"][r][c] if c < len(fila) else None
                color = RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else COLOR_TEXTO
                _parrafos_en(celda.text_frame, self.par[ids] if ids is not None else [], pt, color)
                if r == 0:
                    for p in celda.text_frame.paragraphs:
                        for run in p.runs:
                            run.font.bold = True
        return fh * len(filas) + 91440


# --------------------------------------------------------------------------
# Presentación
# --------------------------------------------------------------------------
def _registrar(pzs: list[dict], textos: list[str]):
    """Asigna a cada pieza el índice de sus textos en la llamada única a pandoc."""
    for pz in pzs:
        if pz["t"] == "texto":
            pz["id"] = len(textos)
            textos.append(pz["latex"])
        elif pz["t"] == "bloque":
            pz["id_titulo"] = len(textos)
            textos.append(pz["titulo"] or " ")
            _registrar(pz["hijos"], textos)
        elif pz["t"] == "columnas":
            for col in pz["cols"]:
                _registrar(col, textos)
        elif pz["t"] == "tabla":
            pz["ids"] = []
            for fila in pz["filas"]:
                pz["ids"].append([])
                for celda in fila:
                    pz["ids"][-1].append(len(textos))
                    textos.append(celda or " ")


def _slide_num_xml(prs):
    for sh in prs.slide_master.shapes:
        if sh.is_placeholder and "Number" in sh.name or "número" in sh.name.lower():
            return copy.deepcopy(sh._element)
    return None


def _nueva(prs, sldnum):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    if sldnum is not None:
        el = copy.deepcopy(sldnum)
        txb = el.find(qn("p:txBody"))
        if txb is not None:                           # el número como campo de diapositiva
            for p in txb.findall(qn("a:p")):
                txb.remove(p)
            p = etree.SubElement(txb, qn("a:p"))
            fld = etree.SubElement(p, qn("a:fld"))
            fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
            fld.set("type", "slidenum")
            etree.SubElement(fld, qn("a:t")).text = "‹#›"
        s.shapes._spTree.append(el)
    return s


def _titulo(slide, parrafos, texto_plano: str, fuentes: dict | None = None, latex: str = "") -> int:
    """Dibuja el título y devuelve dónde puede empezar el cuerpo (EMU)."""
    x, y, w, h = TITULO_XY
    pt = 28 if len(texto_plano) <= 34 else 24 if len(texto_plano) <= 44 else 20
    lineas = max(1, -(-len(texto_plano) // int(w / (pt * EMU_PT * 0.56))))
    h = int(lineas * pt * EMU_PT * 1.25) + 91440
    caja, tf = _caja(slide, x, y, w, h)
    if fuentes is not None:
        fuentes[caja.shape_id] = (latex, pt, COLOR_TITULO, False)
    _parrafos_en(tf, parrafos, pt, COLOR_TITULO)
    for r in tf._txBody.iter(qn("a:rPr")):
        r.set("spc", "100")
        lat = r.find(qn("a:latin"))
        if lat is None:
            lat = etree.SubElement(r, qn("a:latin"))
        lat.set("typeface", FUENTE_TITULO)
    return max(CUERPO[1], y + h + 45720)


def exportar(frames: list[tuple[str, str]], outline: dict, destino: Path, *,
             plantilla: Path | None = None, figuras: Path | None = None,
             pdf_beamer: Path | None = None, macros: str = "", notas: list[str] | None = None) -> Path:
    """frames: [(kind, frame_latex)] en orden (incluye portada y agenda). notas: la del expositor
    de cada frame (mismo orden; "" si no tiene), va al panel de notas. Devuelve destino."""
    notas = notas or [""] * len(frames)
    prs = Presentation(str(plantilla or DEFAULT_PLANTILLA))
    sldnum = _slide_num_xml(prs)
    contenido = [(k, f) for k, f in frames if k not in ("title", "agenda")]
    notas_contenido = [n for (k, _), n in zip(frames, notas) if k not in ("title", "agenda")]

    # 1. piezas de cada frame y una sola llamada a pandoc para todo el texto
    textos: list[str] = []
    plan = []
    for kind, frame in contenido:
        titulo, cuerpo = partes_frame(frame)
        pzs = piezas(cuerpo)
        idt = len(textos)
        textos.append(titulo or " ")
        _registrar(pzs, textos)
        plan.append((kind, titulo, idt, pzs, frame))
    meta = [outline.get("title", ""), outline.get("authors", ""), outline.get("venue", "")]
    base_meta = len(textos)
    textos += [m or " " for m in meta]
    secciones = list(dict.fromkeys(s.get("section") for s in outline.get("slides", []) if s.get("section")))
    base_sec = len(textos)
    textos += secciones
    par = dict(enumerate(a_parrafos(textos, macros)))

    # 2. portada: la primera diapositiva de la plantilla, con sus textos reemplazados
    if prs.slides:
        portada = prs.slides[0]
        for sh in portada.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                tf = sh.text_frame
                rpr = tf.paragraphs[0].runs[0]._r.find(qn("a:rPr")) if tf.paragraphs[0].runs else None
                txb = tf._txBody
                for p in txb.findall(qn("a:p")):
                    txb.remove(p)
                for n, (k, pt) in enumerate(((0, 26), (1, 16), (2, 14))):
                    for p in par[base_meta + k]:
                        p = copy.deepcopy(p)
                        for r in p.iter(qn("a:rPr")):
                            if rpr is not None:
                                for a, v in rpr.attrib.items():
                                    r.set(a, v)
                                for hijo in rpr:
                                    r.append(copy.deepcopy(hijo))
                            r.set("sz", str(pt * 100))
                            r.set("b", "1" if k == 0 else "0")
                        txb.append(p)
                break
        nota_portada = next((n for (k, _), n in zip(frames, notas) if k == "title"), "")
        if nota_portada:
            _nota(portada, nota_portada)

    # 3. agenda
    rast = _Raster(pdf_beamer)
    if any(k == "agenda" for k, _ in frames) and secciones:
        s = _nueva(prs, sldnum)
        _titulo(s, par[base_meta][:0] or [], "Contenido")
        _set_texto_simple(s.shapes[-1].text_frame, "Contenido", 28, COLOR_TITULO)
        x, y, w, h = CUERPO
        _, tf = _caja(s, x, y, w, h)
        _parrafos_en(tf, [p for n in range(len(secciones)) for p in par[base_sec + n]], 24, COLOR_TEXTO)
        for p in tf._txBody.findall(qn("a:p")):
            _pintar(_vineta(p), 24, COLOR_TEXTO)

    # 4. diapositivas de contenido
    pendientes = []
    for (kind, titulo, idt, pzs, frame), nota in zip(plan, notas_contenido):
        s = _nueva(prs, sldnum)
        if nota:
            _nota(s, nota)
        fuentes: dict = {}
        y0 = _titulo(s, par[idt], _texto_plano(titulo), fuentes, titulo)
        x, _, w, _ = CUERPO
        y, h = y0, CUERPO[1] + CUERPO[3] - y0
        dib = Dibujante(prs, par, figuras, lambda pz, ti=titulo: rast.recorte(ti), fuentes)
        fijo = sum(_alto(p, w, 12) for p in pzs if p["t"] in ("imagen", "raster"))
        pt = next((t for t in TAMANOS if sum(_alto(p, w, t) for p in pzs
                                             if p["t"] not in ("imagen", "raster")) + fijo <= h), TAMANOS[-1])
        for k, pz in enumerate(pzs):
            libre = h - (y - y0)
            resto = sum(_alto(q, w, pt) for q in pzs[k + 1:] if q["t"] not in ("imagen", "raster"))
            y += dib.pieza(s, pz, x, y, w, pt, libre, resto)
        pendientes.append((s, fuentes))
    _fallbacks(pendientes, macros)
    rast.cerrar()
    destino.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(destino))
    return destino


def _nota(slide, texto: str) -> None:
    """Notas del expositor en el panel de notas (un párrafo por párrafo del texto)."""
    tf = slide.notes_slide.notes_text_frame
    parrafos = [p.strip() for p in texto.split("\n\n") if p.strip()] or [texto]
    tf.text = parrafos[0]
    for t in parrafos[1:]:
        tf.add_paragraph().text = t


MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"


def _envolver_matematica(slide, fuentes: dict, imagenes: dict):
    """Las cajas y tablas con ecuaciones (a14:m) van en mc:AlternateContent, como las escribe
    PowerPoint: la opción a14 es la ecuación editable; el fallback (lo que muestran LibreOffice
    o Google Slides) es una imagen del mismo texto compilado con LaTeX o, en tablas, texto."""
    arbol = slide.shapes._spTree
    for el in list(arbol):
        if el.tag not in (qn("p:sp"), qn("p:graphicFrame")) or el.find(".//{%s}m" % A14) is None:
            continue
        fallback = None
        png = imagenes.get(_sid(el))
        if el.tag == qn("p:sp") and png:
            off = el.find(".//" + qn("a:off"))
            x, y = int(off.get("x")), int(off.get("y"))
            ancho, alto = png[1]
            pic = slide.shapes.add_picture(str(png[0]), x, y + 45720, ancho, alto)
            fallback = pic._element
            arbol.remove(fallback)
        elif el.tag == qn("p:graphicFrame"):
            fallback = copy.deepcopy(el)
            for m in list(fallback.iter("{%s}m" % A14)):
                r = etree.Element(qn("a:r"))
                etree.SubElement(r, qn("a:rPr")).set("sz", "1200")
                etree.SubElement(r, qn("a:t")).text = _omml_a_texto(m)
                m.getparent().replace(m, r)
        alt = etree.Element("{%s}AlternateContent" % MC, nsmap={"mc": MC})
        choice = etree.SubElement(alt, "{%s}Choice" % MC, nsmap={"a14": A14})
        choice.set("Requires", "a14")
        arbol.replace(el, alt)
        choice.append(el)
        if fallback is not None:
            etree.SubElement(alt, "{%s}Fallback" % MC).append(fallback)


def _sid(el) -> int:
    """Id de la forma dentro de la diapositiva (estable, a diferencia de id() en lxml)."""
    c = el.find(".//" + qn("p:cNvPr"))
    return int(c.get("id")) if c is not None else -1


def _omml_a_texto(m) -> str:
    return "".join(t.text or "" for t in m.iter("{http://schemas.openxmlformats.org/officeDocument/2006/math}t"))


def _fallbacks(pendientes: list, macros: str):
    """Compila en paralelo las imágenes de fallback de las cajas con ecuaciones."""
    from concurrent.futures import ThreadPoolExecutor
    trabajos = []
    for s, fuentes in pendientes:
        for el in s.shapes._spTree:
            if el.tag == qn("p:sp") and el.find(".//{%s}m" % A14) is not None and _sid(el) in fuentes:
                ext = el.find(".//" + qn("a:ext"))
                trabajos.append((s, el, fuentes[_sid(el)], int(ext.get("cx"))))
    tmp = Path(tempfile.mkdtemp())
    with ThreadPoolExecutor(max_workers=4) as ex:
        res = list(ex.map(lambda t: render_latex(t[2][0], t[2][1], t[2][2], t[3], macros, tmp, t[2][3]),
                          trabajos))
    por_slide: dict = {}
    for (s, el, _, _), png in zip(trabajos, res):
        if png:
            por_slide.setdefault(id(s), {})[_sid(el)] = png
    for s, fuentes in pendientes:
        _envolver_matematica(s, fuentes, por_slide.get(id(s), {}))


def render_latex(latex: str, pt: int, color: RGBColor, ancho_emu: int, macros: str, tmp: Path,
                 negrita: bool = False):
    """Compila un trozo de LaTeX al tamaño y color de la caja; (png, (ancho, alto) EMU) o None."""
    import hashlib
    import pymupdf
    ancho_pt = max(40, ancho_emu / EMU_PT - 8)
    cuerpo = re.sub(r"\\(small|footnotesize|scriptsize|tiny|normalsize)\b", "", latex)
    doc = (f"\\documentclass[varwidth={ancho_pt:.0f}pt,border=1pt]{{standalone}}\n"
           "\\usepackage[utf8]{inputenc}\\usepackage[T1]{fontenc}\\usepackage{amsmath,amssymb,bm,xcolor}\n"
           "\\usepackage{helvet}\\renewcommand{\\familydefault}{\\sfdefault}\n"
           f"{macros}\n\\providecommand{{\\alert}}[1]{{\\textbf{{\\color[HTML]{{{COLOR_ALERT}}}#1}}}}\n"
           "\\begin{document}\n"
           f"\\fontsize{{{pt}}}{{{pt * 1.25:.1f}}}\\selectfont\\color[HTML]{{{color}}}\\raggedright"
           + ("\\bfseries " if negrita else "") + f"\n{cuerpo}\n\\end{{document}}\n")
    h = hashlib.sha1(doc.encode()).hexdigest()[:12]
    d = tmp / h
    d.mkdir(exist_ok=True)
    (d / "f.tex").write_text(doc)
    try:
        subprocess.run(["pdflatex", "-no-shell-escape", "-interaction=nonstopmode", "f.tex"], cwd=d,
                       capture_output=True, timeout=60)
    except (subprocess.TimeoutExpired, FileNotFoundError):
        return None
    pdf = d / "f.pdf"
    if not pdf.exists():
        return None
    with pymupdf.open(str(pdf)) as f:
        pg = f[0]
        png = d / "f.png"
        pg.get_pixmap(dpi=300, alpha=True).save(str(png))
        return png, (int(pg.rect.width * EMU_PT), int(pg.rect.height * EMU_PT))


def _vineta(p):
    ppr = p.find(qn("a:pPr"))
    if ppr is None:
        ppr = etree.Element(qn("a:pPr"))
        p.insert(0, ppr)
    for b in list(ppr):
        if b.tag == qn("a:buNone"):
            ppr.remove(b)
    ppr.set("lvl", "0")
    return p


def _set_texto_simple(tf, texto, pt, color):
    tf.text = texto
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size, r.font.color.rgb, r.font.name = Pt(pt), color, FUENTE_TITULO


class _Raster:
    """Recorta de la presentación Beamer compilada el cuerpo de una diapo (bajo el título)."""
    def __init__(self, pdf: Path | None):
        self.doc = None
        self.tmp = Path(tempfile.mkdtemp())
        if pdf and Path(pdf).exists():
            import pymupdf
            self.doc = pymupdf.open(str(pdf))

    def recorte(self, titulo: str):
        if self.doc is None:
            return None
        import pymupdf
        import unicodedata
        # solo letras y dígitos, sin acentos: el PDF puede extraer «ó» como «o´» o ligaduras
        # (y las ligaduras fi/fl/ff pueden desaparecer: se ignoran esas letras al comparar)
        def norm(t: str) -> str:
            return re.sub(r"[^a-z0-9]|[fil]", "", "".join(
                c for c in unicodedata.normalize("NFKD", t.lower()) if not unicodedata.combining(c)))
        clave = norm(_texto_plano(titulo))[:30]
        for n, pg in enumerate(self.doc):
            if clave and clave in norm(pg.get_text()):
                r = pg.rect
                zona = pymupdf.Rect(r.x0, r.y0 + r.height * 0.14, r.x1, r.y1 - r.height * 0.04)
                # algorithm2e (ruled) dibuja reglas horizontales arriba y abajo del algoritmo
                reglas = [d["rect"] for d in pg.get_drawings() if zona.contains(d["rect"].tl)
                          and d["rect"].height < 2 and d["rect"].width > r.width * 0.2]
                if len(reglas) >= 2:
                    clip = pymupdf.Rect(min(c.x0 for c in reglas), min(c.y0 for c in reglas),
                                        max(c.x1 for c in reglas), max(c.y1 for c in reglas))
                    clip = clip + (-3, -3, 3, 3)
                else:
                    cajas = [pymupdf.Rect(b[:4]) for b in pg.get_text("blocks")
                             if pymupdf.Rect(b[:4]).intersects(zona)]
                    clip = zona
                    if cajas:
                        clip = cajas[0]
                        for c in cajas[1:]:
                            clip |= c
                        clip = (clip & zona) + (-4, -4, 4, 4)
                f = self.tmp / f"r{n}.png"
                pg.get_pixmap(dpi=220, clip=clip).save(str(f))
                return f
        return None

    def cerrar(self):
        if self.doc is not None:
            self.doc.close()
